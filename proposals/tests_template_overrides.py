"""
Admin-managed replacements for the bundled document templates.

The proposal flow generates its DOCX/XLSX downloads from templates that ship
under ``proposals/template_files/``. Administrators may replace any of them
through the admin screens; these tests pin down that behaviour:

* the resolver prefers the admin upload and falls back to the bundled copy
* generated documents and downloads actually use the replacement
* uploads are validated (right slot, right file type) and can be reset
* the admin screens stay admin-only
"""

import io
import zipfile

from django.contrib import admin as django_admin
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from accounts.models import Profile
from accounts.tests import factories
from openpyxl import load_workbook

from .models import CustomProposalTemplate, Proposal, ProposalTemplateOverride
from .template_store import (
    TEMPLATE_FILES,
    TemplateNotFound,
    bundled_path,
    is_valid_replacement,
    open_template,
)
from .docx_forms import build_extension_form_docx


def docx_bytes_with_marker(marker):
    """A real .docx carrying one paragraph with ``marker``."""
    from docx import Document

    document = Document()
    document.add_paragraph(marker)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes_with_marker_sheet(marker):
    """A real .xlsx whose first sheet is named ``marker``."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.active.title = marker[:31]
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def response_body(response):
    """The bytes a response sent, streaming (FileResponse) or not."""
    if getattr(response, "streaming", False):
        return b"".join(response.streaming_content)
    return response.content



class TemplateStoreTests(TestCase):
    """The resolver: override wins, bundled copy is the fallback."""

    def test_every_registered_slot_has_a_bundled_file(self):
        for key in TEMPLATE_FILES:
            with self.subTest(key=key):
                self.assertTrue(bundled_path(key).exists(), f"{key} is missing")

    def test_without_an_override_the_bundled_file_is_used(self):
        stream, filename, is_override = open_template("form2_training_design_template.docx")
        try:
            self.assertFalse(is_override)
            self.assertEqual(filename, "form2_training_design_template.docx")
            self.assertEqual(stream.read(), bundled_path("form2_training_design_template.docx").read_bytes())
        finally:
            stream.close()

    def test_an_override_replaces_the_bundled_file(self):
        marker = docx_bytes_with_marker("OVERRIDE-RESOLVER-MARKER")
        ProposalTemplateOverride.objects.create(
            key="form2_training_design_template.docx",
            file=SimpleUploadedFile("replacement.docx", marker),
        )

        stream, filename, is_override = open_template("form2_training_design_template.docx")
        try:
            self.assertTrue(is_override)
            self.assertEqual(stream.read(), marker)
        finally:
            stream.close()

    def test_a_missing_key_is_rejected(self):
        with self.assertRaises(ValueError):
            open_template("not_a_template.docx")

    def test_validation_requires_the_slot_extension(self):
        self.assertTrue(is_valid_replacement("form1_project_template.docx", "anything.DOCX"))
        self.assertFalse(is_valid_replacement("form1_project_template.docx", "anything.xlsx"))
        self.assertFalse(is_valid_replacement("project_work_plan_template.xlsx", "anything.docx"))
        self.assertFalse(is_valid_replacement("unknown_key.docx", "anything.docx"))


class GeneratedDocumentOverrideTests(TestCase):
    """Generated downloads must use the admin replacement."""

    @classmethod
    def setUpTestData(cls):
        cls.owner = factories.make_user("tpl_owner", Profile.ROLE_FACULTY)

    def setUp(self):
        self.proposal = Proposal.objects.create(
            created_by=self.owner,
            title="Template Override Proposal",
            extension_type="COMMUNITY_BASED",
            scope_type="PROJECT",
        )

    def test_the_docx_generator_uses_the_override(self):
        marker = docx_bytes_with_marker("OVERRIDE-DOCX-MARKER")
        ProposalTemplateOverride.objects.create(
            key="form2_training_design_template.docx",
            file=SimpleUploadedFile("replacement.docx", marker),
        )

        data = build_extension_form_docx(proposal=self.proposal)

        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
        self.assertIn("OVERRIDE-DOCX-MARKER", xml)

    def test_the_work_plan_download_uses_the_override(self):
        marker = xlsx_bytes_with_marker_sheet("OVERRIDE-MARKER-SHEET")
        ProposalTemplateOverride.objects.create(
            key="project_work_plan_template.xlsx",
            file=SimpleUploadedFile("replacement.xlsx", marker),
        )

        client = factories.make_client(self.owner)
        response = client.get(reverse("download_work_plan_template", args=[self.proposal.id]))

        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(io.BytesIO(response.content))
        self.assertIn("OVERRIDE-MARKER-SHEET", workbook.sheetnames)

    def test_downloads_still_work_without_any_override(self):
        client = factories.make_client(self.owner)
        for name in (
            "download_work_plan_template",
            "download_gantt_chart_template",
            "download_funding_template",
        ):
            with self.subTest(view=name):
                response = client.get(reverse(name, args=[self.proposal.id]))
                self.assertEqual(
                    response.status_code,
                    200,
                    f"{name} should serve the bundled template",
                )


class AdminTemplateScreensTests(TestCase):
    """The 'Proposal Templates' screen: admin-only, validated, resettable."""

    @classmethod
    def setUpTestData(cls):
        cls.admin_user, cls.admin_client = factories.admin()
        cls.faculty = factories.make_user("tpl_faculty", Profile.ROLE_FACULTY)

    def test_only_admins_can_open_the_screen(self):
        self.assertEqual(
            self.admin_client.get(reverse("proposal_templates_list")).status_code, 200
        )
        faculty_status = factories.make_client(self.faculty).get(
            reverse("proposal_templates_list")
        ).status_code
        self.assertIn(faculty_status, (302, 403))

    def test_the_screen_lists_every_template_slot(self):
        html = self.admin_client.get(reverse("proposal_templates_list")).content.decode()
        for key, (label, _ext) in TEMPLATE_FILES.items():
            with self.subTest(key=key):
                self.assertIn(label, html)

    def test_a_valid_upload_replaces_the_template(self):
        marker = docx_bytes_with_marker("UPLOADED-BY-ADMIN")
        response = self.admin_client.post(
            reverse("proposal_template_replace"),
            {
                "key": "form1_project_template.docx",
                "file": SimpleUploadedFile("new_form1.docx", marker),
                "notes": "2026 revision",
            },
        )
        self.assertEqual(response.status_code, 302)

        override = ProposalTemplateOverride.objects.get(key="form1_project_template.docx")
        self.assertEqual(override.notes, "2026 revision")
        self.assertEqual(override.uploaded_by, self.admin_user)

    def test_a_wrong_extension_is_refused(self):
        response = self.admin_client.post(
            reverse("proposal_template_replace"),
            {
                "key": "project_work_plan_template.xlsx",
                "file": SimpleUploadedFile("not-a-spreadsheet.docx", docx_bytes_with_marker("X")),
            },
        )
        self.assertEqual(response.status_code, 302)  # back to the list, with an error
        self.assertFalse(
            ProposalTemplateOverride.objects.filter(key="project_work_plan_template.xlsx").exists()
        )

    def test_a_non_office_file_is_refused(self):
        response = self.admin_client.post(
            reverse("proposal_template_replace"),
            {
                "key": "form1_project_template.docx",
                "file": SimpleUploadedFile("plain.docx", b"this is not a zip archive"),
            },
        )
        self.assertEqual(response.status_code, 302)
        self.assertFalse(
            ProposalTemplateOverride.objects.filter(key="form1_project_template.docx").exists()
        )

    def test_reset_restores_the_bundled_template(self):
        ProposalTemplateOverride.objects.create(
            key="clear_summary_template.docx",
            file=SimpleUploadedFile("replacement.docx", docx_bytes_with_marker("X")),
        )

        response = self.admin_client.post(
            reverse("proposal_template_reset"), {"key": "clear_summary_template.docx"}
        )

        self.assertEqual(response.status_code, 302)
        self.assertFalse(
            ProposalTemplateOverride.objects.filter(key="clear_summary_template.docx").exists()
        )

    def test_reset_rejects_unknown_slots(self):
        response = self.admin_client.post(
            reverse("proposal_template_reset"), {"key": "not_a_template.docx"}
        )
        self.assertEqual(response.status_code, 302)

    def test_replacing_twice_updates_the_same_slot(self):
        first = docx_bytes_with_marker("FIRST")
        second = docx_bytes_with_marker("SECOND")
        for payload in (first, second):
            self.admin_client.post(
                reverse("proposal_template_replace"),
                {
                    "key": "form2_training_design_template.docx",
                    "file": SimpleUploadedFile("r.docx", payload),
                },
            )

        self.assertEqual(
            ProposalTemplateOverride.objects.filter(key="form2_training_design_template.docx").count(),
            1,
        )
        stream, _name, is_override = open_template("form2_training_design_template.docx")
        try:
            self.assertTrue(is_override)
            self.assertEqual(stream.read(), second)
        finally:
            stream.close()


class DjangoAdminRegistrationTests(TestCase):
    """The overrides are also editable through the built-in Django admin."""

    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth import get_user_model

        cls.superuser = get_user_model().objects.create_superuser(
            username="nx_superuser", email="su@example.com", password="sup3r-Secret"
        )

    def test_the_model_is_registered(self):
        self.assertIn(ProposalTemplateOverride, django_admin.site._registry)

    def test_the_changelist_and_add_form_are_reachable(self):
        client = factories.make_client(self.superuser)
        changelist = client.get("/admin/proposals/proposaltemplateoverride/")
        self.assertEqual(changelist.status_code, 200)
        add_form = client.get("/admin/proposals/proposaltemplateoverride/add/")
        self.assertEqual(add_form.status_code, 200)


class TemplateDownloadTests(TestCase):
    """Every template on the screen can be downloaded, default included."""

    @classmethod
    def setUpTestData(cls):
        cls.admin_user, cls.admin_client = factories.admin("tpl_dl_admin")
        cls.faculty = factories.make_user("tpl_dl_faculty", Profile.ROLE_FACULTY)

    def test_a_slot_without_an_override_downloads_the_bundled_file(self):
        key = "project_work_plan_template.xlsx"
        response = self.admin_client.get(reverse("proposal_template_download", args=[key]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertEqual(response_body(response), bundled_path(key).read_bytes())

    def test_a_replaced_slot_downloads_the_uploaded_file(self):
        marker = docx_bytes_with_marker("DOWNLOAD-THE-OVERRIDE")
        key = "form2_training_design_template.docx"
        ProposalTemplateOverride.objects.create(
            key=key,
            file=SimpleUploadedFile("revised_form2.docx", marker),
        )

        response = self.admin_client.get(reverse("proposal_template_download", args=[key]))

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertEqual(response_body(response), marker)

    def test_an_unknown_slot_is_not_found(self):
        response = self.admin_client.get(
            reverse("proposal_template_download", args=["not_a_template.docx"])
        )
        self.assertEqual(response.status_code, 404)

    def test_only_admins_can_download(self):
        status = factories.make_client(self.faculty).get(
            reverse("proposal_template_download", args=["clear_summary_template.docx"])
        ).status_code
        self.assertIn(status, (302, 403))


class BuiltInTemplateEditTests(TestCase):
    """The per-slot edit screen: note, file, or both."""

    @classmethod
    def setUpTestData(cls):
        cls.admin_user, cls.admin_client = factories.admin("tpl_edit_admin")
        cls.faculty = factories.make_user("tpl_edit_faculty", Profile.ROLE_FACULTY)

    def setUp(self):
        self.key = "form2_training_design_template.docx"
        self.url = reverse("proposal_template_edit", args=[self.key])

    def test_the_edit_screen_opens_for_every_slot(self):
        for key in TEMPLATE_FILES:
            with self.subTest(key=key):
                response = self.admin_client.get(reverse("proposal_template_edit", args=[key]))
                self.assertEqual(response.status_code, 200)
                self.assertIn(TEMPLATE_FILES[key][0], response.content.decode())

    def test_an_unknown_slot_is_not_found(self):
        response = self.admin_client.get(
            reverse("proposal_template_edit", args=["not_a_template.docx"])
        )
        self.assertEqual(response.status_code, 404)

    def test_only_admins_can_edit(self):
        status = factories.make_client(self.faculty).get(self.url).status_code
        self.assertIn(status, (302, 403))

    def test_editing_the_note_keeps_the_file(self):
        original = docx_bytes_with_marker("EDIT-KEEPS-FILE")
        ProposalTemplateOverride.objects.create(
            key=self.key,
            file=SimpleUploadedFile("current.docx", original),
            notes="old note",
        )

        response = self.admin_client.post(self.url, {"notes": "2026 budget revision"})

        self.assertEqual(response.status_code, 302)
        override = ProposalTemplateOverride.objects.get(key=self.key)
        self.assertEqual(override.notes, "2026 budget revision")
        stream, _name, is_override = open_template(self.key)
        try:
            self.assertTrue(is_override)
            self.assertEqual(stream.read(), original)
        finally:
            stream.close()

    def test_a_new_file_replaces_the_stored_one(self):
        ProposalTemplateOverride.objects.create(
            key=self.key,
            file=SimpleUploadedFile("current.docx", docx_bytes_with_marker("OLD")),
        )
        revised = docx_bytes_with_marker("REVISED-BY-ADMIN")

        response = self.admin_client.post(
            self.url,
            {"notes": "new format", "file": SimpleUploadedFile("revised.docx", revised)},
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(ProposalTemplateOverride.objects.filter(key=self.key).count(), 1)
        stream, _name, _is_override = open_template(self.key)
        try:
            self.assertEqual(stream.read(), revised)
        finally:
            stream.close()

    def test_a_slot_without_a_replacement_needs_a_file(self):
        response = self.admin_client.post(self.url, {"notes": "no file yet"})

        self.assertEqual(response.status_code, 200)  # re-rendered with the error
        self.assertFalse(ProposalTemplateOverride.objects.filter(key=self.key).exists())

    def test_a_wrong_extension_is_refused(self):
        response = self.admin_client.post(
            self.url,
            {"file": SimpleUploadedFile("sheet.xlsx", xlsx_bytes_with_marker_sheet("X"))},
        )

        self.assertEqual(response.status_code, 200)
        self.assertFalse(ProposalTemplateOverride.objects.filter(key=self.key).exists())


class CustomProposalTemplateTests(TestCase):
    """Files the office adds for a process or format that is not wired in yet."""

    @classmethod
    def setUpTestData(cls):
        cls.admin_user, cls.admin_client = factories.admin("tpl_custom_admin")
        cls.faculty = factories.make_user("tpl_custom_faculty", Profile.ROLE_FACULTY)

    def add_file(self, name="MOA_Renewal_Form.docx", content=None, **extra):
        payload = {
            "files": SimpleUploadedFile(
                name, content if content is not None else docx_bytes_with_marker("CUSTOM")
            ),
            "notes": extra.pop("notes", "for the 2027 process"),
        }
        payload.update(extra)
        return self.admin_client.post(reverse("proposal_custom_template_add"), payload)

    def test_an_admin_can_add_a_file(self):
        response = self.add_file()

        self.assertEqual(response.status_code, 302)
        template = CustomProposalTemplate.objects.get()
        self.assertEqual(template.title, "MOA Renewal Form")
        self.assertEqual(template.key, "moa-renewal-form.docx")
        self.assertEqual(template.uploaded_by, self.admin_user)
        self.assertTrue(template.is_active)

    def test_several_files_can_be_added_at_once(self):
        response = self.admin_client.post(
            reverse("proposal_custom_template_add"),
            {
                "files": [
                    SimpleUploadedFile("Evaluation_Form.docx", docx_bytes_with_marker("A")),
                    SimpleUploadedFile("Budget_Matrix.xlsx", xlsx_bytes_with_marker_sheet("B")),
                ],
                "notes": "incoming formats",
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(
            sorted(CustomProposalTemplate.objects.values_list("key", flat=True)),
            ["budget-matrix.xlsx", "evaluation-form.docx"],
        )

    def test_the_same_title_gets_a_second_key(self):
        self.add_file()
        self.add_file()

        self.assertEqual(
            sorted(CustomProposalTemplate.objects.values_list("key", flat=True)),
            ["moa-renewal-form-2.docx", "moa-renewal-form.docx"],
        )

    def test_a_key_never_collides_with_a_built_in_slot(self):
        response = self.add_file(name="Clearance_Summary.docx")

        self.assertEqual(response.status_code, 302)
        self.assertEqual(CustomProposalTemplate.objects.get().key, "clearance-summary.docx")
        self.assertIn("clearance-summary.docx", [t.key for t in CustomProposalTemplate.objects.all()])
        self.assertNotIn("clearance-summary.docx", TEMPLATE_FILES)

    def test_an_unsupported_format_is_refused(self):
        response = self.add_file(name="photo.png", content=b"\x89PNG\r\n\x1a\n" + b"0" * 32)

        self.assertEqual(response.status_code, 302)
        self.assertFalse(CustomProposalTemplate.objects.exists())

    def test_a_renamed_file_is_refused(self):
        response = self.add_file(name="pretend.docx", content=b"this is not a zip archive")

        self.assertEqual(response.status_code, 302)
        self.assertFalse(CustomProposalTemplate.objects.exists())

    def test_added_files_appear_on_the_screen(self):
        self.add_file()
        html = self.admin_client.get(reverse("proposal_templates_list")).content.decode()

        self.assertIn("MOA Renewal Form", html)
        self.assertIn("moa-renewal-form.docx", html)
        self.assertIn("Added template files", html)

    def test_an_added_file_can_be_downloaded(self):
        marker = docx_bytes_with_marker("CUSTOM-DOWNLOAD")
        self.add_file(content=marker)
        template = CustomProposalTemplate.objects.get()

        response = self.admin_client.get(
            reverse("proposal_custom_template_download", args=[template.pk])
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertEqual(response_body(response), marker)

    def test_an_added_file_can_be_edited(self):
        self.add_file()
        template = CustomProposalTemplate.objects.get()
        revised = docx_bytes_with_marker("CUSTOM-REVISED")

        response = self.admin_client.post(
            reverse("proposal_custom_template_edit", args=[template.pk]),
            {
                "title": "MOA Renewal Form 2027",
                "notes": "updated for the new signatories",
                "file": SimpleUploadedFile("renewal_2027.docx", revised),
                "is_active": "on",
            },
        )

        self.assertEqual(response.status_code, 302)
        template.refresh_from_db()
        self.assertEqual(template.title, "MOA Renewal Form 2027")
        self.assertEqual(template.notes, "updated for the new signatories")
        self.assertEqual(template.key, "moa-renewal-form.docx")  # identity is stable
        self.assertTrue(template.is_active)
        template.file.open("rb")
        self.assertEqual(template.file.read(), revised)

    def test_an_edit_can_retire_a_file_without_touching_it(self):
        self.add_file()
        template = CustomProposalTemplate.objects.get()

        response = self.admin_client.post(
            reverse("proposal_custom_template_edit", args=[template.pk]),
            {"title": template.title, "notes": "superseded"},
        )

        self.assertEqual(response.status_code, 302)
        template.refresh_from_db()
        self.assertFalse(template.is_active)
        self.assertTrue(template.file)

    def test_an_added_file_can_be_deleted(self):
        self.add_file()
        template = CustomProposalTemplate.objects.get()

        response = self.admin_client.post(
            reverse("proposal_custom_template_delete", args=[template.pk])
        )

        self.assertEqual(response.status_code, 302)
        self.assertFalse(CustomProposalTemplate.objects.exists())

    def test_only_admins_can_manage_added_files(self):
        self.add_file()
        template = CustomProposalTemplate.objects.get()
        client = factories.make_client(self.faculty)

        for url in (
            reverse("proposal_custom_template_edit", args=[template.pk]),
            reverse("proposal_custom_template_download", args=[template.pk]),
        ):
            with self.subTest(url=url):
                self.assertIn(client.get(url).status_code, (302, 403))

        self.assertIn(
            client.post(reverse("proposal_custom_template_add"), {}).status_code, (302, 403)
        )
        self.assertIn(
            client.post(
                reverse("proposal_custom_template_delete", args=[template.pk])
            ).status_code,
            (302, 403),
        )
        self.assertTrue(CustomProposalTemplate.objects.filter(pk=template.pk).exists())


class CustomTemplateAdminRegistrationTests(TestCase):
    """Added files are editable through the built-in Django admin too."""

    @classmethod
    def setUpTestData(cls):
        from django.contrib.auth import get_user_model

        cls.superuser = get_user_model().objects.create_superuser(
            username="nx_custom_su", email="csu@example.com", password="sup3r-Secret"
        )

    def test_the_model_is_registered(self):
        self.assertIn(CustomProposalTemplate, django_admin.site._registry)

    def test_the_changelist_and_add_form_are_reachable(self):
        client = factories.make_client(self.superuser)
        self.assertEqual(
            client.get("/admin/proposals/customproposaltemplate/").status_code, 200
        )
        self.assertEqual(
            client.get("/admin/proposals/customproposaltemplate/add/").status_code, 200
        )
